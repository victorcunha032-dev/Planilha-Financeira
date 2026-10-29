import openpyxl, json, re, sys
wb=openpyxl.load_workbook(sys.argv[1])
def split(s):
    s=s.strip()
    m=re.match(r'^(.*?)\s*([^\w\s].*)?$',s)
    parts=s.rsplit(' ',1)
    if len(parts)==2 and not re.search(r'\w',parts[1]): return parts[0].strip(),parts[1]
    return s,''
ws=wb['Categorias']
cats=[]
for r in range(7,25):
    e=ws.cell(r,5).value
    if not e or e.startswith('Despesas'): continue
    tipo='R' if r<15 else 'D'
    n,i=split(e); cats.append({'tipo':tipo,'nome':n,'icone':i,'desc':ws.cell(r,6).value or '','subs':[]})
cats.append({'tipo':'D','nome':'Seguros','icone':'🛡️','desc':'Previdência, seguro de vida e seguro do carro.','subs':[]})
byname={c['nome'].lower():c for c in cats}
cur=None
for r in range(6,80):
    v=ws.cell(r,10).value
    if not v: continue
    v=v.strip()
    if v.lower() in byname: cur=byname[v.lower()]; continue
    if cur and v not in cur['subs']: cur['subs'].append(v)
ws=wb['Central de lançamentos']
saldo=ws['G3'].value
lan=[]
for r in range(6,1004):
    d=ws.cell(r,2).value
    if not d: continue
    g,_=split(ws.cell(r,4).value); sub=ws.cell(r,5).value.strip()
    c=byname[g.lower()]
    if sub not in c['subs']: c['subs'].append(sub)
    lan.append({'id':'l%d'%r,'data':d.strftime('%Y-%m-%d'),'tipo':'R' if ws.cell(r,3).value=='Receitas' else 'D','grupo':c['nome'],'sub':sub,'valor':ws.cell(r,6).value,'desc':''})
ws=wb['Contas a Pagar']
contas=[]
YEAR=2026
for r in range(8,1000):
    n=ws.cell(r,1).value
    if not n: continue
    p=ws.cell(r,2).value
    vals={}
    for m in range(12):
        col=6+2*m
        v=ws.cell(r,col).value
        if isinstance(v,(int,float)):
            vals['%d-%02d'%(YEAR,m+1)]={'v':v,'pago':str(ws.cell(r,col+1).value or '').lower()=='pago'}
    contas.append({'id':'c%d'%r,'nome':n.strip(),'parcelas':p if isinstance(p,(int,float)) else None,'valores':vals})
state={'v':1,'saldoInicial':saldo,'categorias':cats,'lancamentos':lan,'contas':contas}
json.dump(state,open(sys.argv[2],'w'),ensure_ascii=False)
print(json.dumps(cats,ensure_ascii=False)[:1500]); print(len(lan),len(contas),saldo)
